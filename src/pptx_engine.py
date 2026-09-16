import io
from pptx import Presentation

def processar_relatorio(template_path, dados_mapeados):
    prs = Presentation(template_path)

    for slide in prs.slides:
        for shape in slide.shapes:
            # 1. Substituição em caixas de texto
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        for chave, valor in dados_mapeados.items():
                            if chave in run.text:
                                run.text = run.text.replace(chave, str(valor))
            
            # 2. Substituição em tabelas (alinhado com o 'if' de cima)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                for chave, valor in dados_mapeados.items():
                                    if chave in run.text:
                                        run.text = run.text.replace(chave, str(valor))

    pptx_io = io.BytesIO()
    prs.save(pptx_io)
    pptx_io.seek(0)
    return pptx_io
